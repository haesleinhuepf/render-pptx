import io

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from render_pptx import render_presentation, render_slide


def _picture_bytes():
    img = Image.new("RGB", (100, 100), (255, 0, 0))
    img.paste((0, 0, 255), (50, 0, 100, 100))  # right half blue
    buf = io.BytesIO()
    img.save(buf, "PNG")
    buf.seek(0)
    return buf


def _deck(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(2), Inches(2))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0, 255, 0)
    pic = slide.shapes.add_picture(_picture_bytes(), Inches(4), Inches(4), Inches(2), Inches(2))
    pic.crop_left = 0.5  # keep only blue half
    path = tmp_path / "t.pptx"
    prs.save(path)
    return path


def test_shape_and_cropped_picture(tmp_path):
    img = render_slide(_deck(tmp_path), 0, width=1000)
    assert img.size == (1000, 750)
    assert img.getpixel((50, 50)) == (0, 255, 0)
    assert img.getpixel((500, 500)) == (0, 0, 255)  # no red from cropped half
    assert img.getpixel((900, 700)) == (255, 255, 255)


def test_export_png(tmp_path):
    paths = render_presentation(_deck(tmp_path), tmp_path / "out", width=500)
    assert len(paths) == 1 and Image.open(paths[0]).format == "PNG"


def _text_deck(tmp_path, text, rotation=0, tabs=()):
    from lxml import etree
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(8), Inches(1))
    tb.rotation = rotation
    para = tb.text_frame.paragraphs[0]
    para.add_run().text = text
    if tabs:
        ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
        pPr = para._p.get_or_add_pPr()
        lst = etree.SubElement(pPr, "{%s}tabLst" % ns)
        for pos in tabs:
            etree.SubElement(lst, "{%s}tab" % ns, pos=str(int(Inches(pos))), algn="l")
    path = tmp_path / "txt.pptx"
    prs.save(path)
    return path


def _dark_columns(img):
    return [x for x in range(img.width)
            if any(img.getpixel((x, y))[0] < 100 for y in range(img.height))]


def test_tab_uses_ruler_position(tmp_path):
    img = render_slide(_text_deck(tmp_path, "A\tB", tabs=(4,)), 0, width=1000)
    cols = _dark_columns(img)
    # "B" starts at 1in margin + 0.1in inset + 4in tab = 5.1in -> 510px of 1000px (10in)
    assert any(500 <= x <= 530 for x in cols)
    assert not any(130 < x < 500 for x in cols)


def test_text_rotation(tmp_path):
    img = render_slide(_text_deck(tmp_path, "HELLO WORLD", rotation=90), 0, width=1000)
    cols = _dark_columns(img)
    # horizontal text spans wide; rotated by 90 it becomes narrow
    assert max(cols) - min(cols) < 100


def test_indentation_without_tab(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(8), Inches(1))
    para = tb.text_frame.paragraphs[0]
    para.add_run().text = "Indented"
    para._p.get_or_add_pPr().set("marL", str(int(Inches(2))))
    path = tmp_path / "ind.pptx"
    prs.save(path)
    cols = _dark_columns(render_slide(path, 0, width=1000))
    assert min(cols) >= 300  # 1in box + 0.1in inset + 2in indent -> 310px


def test_font_one_point_smaller(tmp_path):
    from render_pptx import renderer
    sizes = []
    orig = renderer._font
    renderer._font = lambda size, bold=False, name=None, italic=False: sizes.append(size) or orig(size, bold, name, italic)
    try:
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        run = tb.text_frame.paragraphs[0].add_run()
        run.text = "x"
        from pptx.util import Pt
        run.font.size = Pt(20)
        path = tmp_path / "f.pptx"
        prs.save(path)
        render_slide(path, 0, width=720)  # 720px / 10in = 72 px per in = 1px per pt
    finally:
        renderer._font = orig
    assert sizes == [19]


def test_inherited_size_and_font_name(tmp_path):
    from render_pptx import renderer
    seen = []
    orig = renderer._font
    renderer._font = lambda size, bold=False, name=None, italic=False: seen.append((size, name)) or orig(size, bold, name, italic)
    try:
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])  # title slide
        slide.shapes.title.text = "T"  # size from master title style (44pt), theme font
        tb = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1))
        run = tb.text_frame.paragraphs[0].add_run()
        run.text = "x"
        run.font.name = "DejaVu Sans"
        path = tmp_path / "i.pptx"
        prs.save(path)
        render_slide(path, 0, width=720)
    finally:
        renderer._font = orig
    assert seen[0][0] > 19  # title is not rendered with the 18pt fallback
    assert seen[0][1] == "Calibri"
    assert seen[-1] == (17, "DejaVu Sans")
