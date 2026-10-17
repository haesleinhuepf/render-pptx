"""Render pptx slides onto a PIL canvas."""
import copy
import io
import math
import os
import re
import subprocess

from PIL import Image, ImageChops, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.dml import MSO_FILL
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu

_FONTS = ["DejaVuSans.ttf", "Arial.ttf", "arial.ttf", "LiberationSans-Regular.ttf"]
_FONTS_BOLD = ["DejaVuSans-Bold.ttf", "Arial Bold.ttf", "arialbd.ttf", "LiberationSans-Bold.ttf"]


_font_files = {}
_installed_fonts = None


def _scan_installed_fonts():
    """Map lower-case 'family [style]' names to font files from the Windows registry and font folders."""
    global _installed_fonts
    if _installed_fonts is not None:
        return _installed_fonts
    found = {}
    try:
        import winreg
        for hive, base in ((winreg.HKEY_LOCAL_MACHINE, r"C:\Windows\Fonts"),
                           (winreg.HKEY_CURRENT_USER,
                            os.path.join(os.environ.get("LOCALAPPDATA", ""), r"Microsoft\Windows\Fonts"))):
            try:
                with winreg.OpenKey(hive, r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts") as key:
                    for i in range(winreg.QueryInfoKey(key)[1]):
                        value, path, _t = winreg.EnumValue(key, i)
                        value = re.sub(r"\s*\((TrueType|OpenType)\)$", "", value).split(" & ")[0]
                        found[value.lower()] = path if os.path.isabs(path) else os.path.join(base, path)
            except OSError:
                continue
    except ImportError:
        pass
    cloud = os.path.join(os.environ.get("LOCALAPPDATA", ""), r"Microsoft\FontCache\4\CloudFonts")
    if os.path.isdir(cloud):
        for root, _dirs, files in os.walk(cloud):
            for fn in files:
                if fn.lower().endswith((".ttf", ".otf")):
                    path = os.path.join(root, fn)
                    try:
                        family, style = ImageFont.truetype(path, 10).getname()
                    except OSError:
                        continue
                    found.setdefault(family.lower() + " " + style.lower(), path)
                    if style.lower() == "regular":
                        found.setdefault(family.lower(), path)
    for folder in [os.environ.get("RENDER_PPTX_FONT_DIR")] + [
            os.path.join(os.path.dirname(__file__), "fonts")]:
        if folder and os.path.isdir(folder):
            for fn in os.listdir(folder):
                stem, ext = os.path.splitext(fn)
                if ext.lower() in (".ttf", ".otf", ".ttc"):
                    found.setdefault(stem.lower().replace("-", " "), os.path.join(folder, fn))
    _installed_fonts = found
    return found


def _font_file_installed(name, bold, italic):
    fonts = _scan_installed_fonts()
    style = " ".join(s for s, on in (("bold", bold), ("italic", italic)) if on)
    base = name.lower()
    for key in ([base + " " + style] if style else []) + ([base] if not style else []) + \
            ([base + " regular"] if not style else []):
        if key in fonts:
            return fonts[key]
    # fall back to the regular face when the requested style is missing
    return fonts.get(base) or fonts.get(base + " regular")


def _font_file(name, bold, italic):
    """Locate a font file for a family name (installed fonts, then fontconfig); None if unavailable."""
    key = (name, bold, italic)
    if key not in _font_files:
        path = _font_file_installed(name, bold, italic)
        if path:
            _font_files[key] = path
            return path
        pattern = name + (":bold" if bold else "") + (":italic" if italic else "")
        try:
            out = subprocess.run(["fc-match", "-f", "%{family}|%{file}", pattern],
                                 capture_output=True, text=True, timeout=10).stdout
            family, _, path = out.partition("|")
            wanted = name.lower()
            if wanted not in [f.strip().lower() for f in family.split(",")]:
                path = None  # fontconfig substituted another family
        except (OSError, subprocess.SubprocessError):
            path = None
        _font_files[key] = path or None
    return _font_files[key]


_SERIF_HINTS = ("serif", "times", "georgia", "garamond", "cambria", "palatino", "book antiqua",
                "bookman", "century", "merriweather", "playfair", "lora", "baskerville", "didot",
                "minion", "caslon", "constantia", "rockwell", "slab")
_MONO_HINTS = ("mono", "code", "courier", "consolas", "menlo", "typewriter", "lucida console")
_SERIF_FALLBACKS = ["Times New Roman", "Liberation Serif", "DejaVu Serif", "Georgia"]
_MONO_FALLBACKS = ["Consolas", "Courier New", "Liberation Mono", "DejaVu Sans Mono"]


def _similar_fonts(name, bold, italic):
    """Candidate font files resembling `name` (by serif / monospace / sans category)."""
    lower = (name or "").lower()
    if any(h in lower for h in _MONO_HINTS):
        families = _MONO_FALLBACKS
    elif any(h in lower for h in _SERIF_HINTS) and "sans" not in lower:
        families = _SERIF_FALLBACKS
    else:
        return []
    return [_font_file(f, bold, italic) for f in families]


def _font(size, bold=False, name=None, italic=False):
    size = max(int(size), 1)
    candidates = []
    if name:
        candidates += [_font_file(name, bold, italic), name + ".ttf"]
        candidates += _similar_fonts(name, bold, italic)
    candidates += (_FONTS_BOLD if bold else []) + _FONTS
    for cand in candidates:
        if not cand:
            continue
        try:
            return ImageFont.truetype(cand, size)
        except OSError:
            continue
    return ImageFont.load_default(size)


_EMOJI_FONTS = ["seguiemj.ttf", "Segoe UI Emoji", "NotoColorEmoji.ttf",
                "/usr/share/fonts/truetype/noto/NotoColorEmoji.ttf",
                "/System/Library/Fonts/Apple Color Emoji.ttc", "Apple Color Emoji.ttc",
                "NotoEmoji-Regular.ttf", "DejaVuSans.ttf"]
_EMOJI_RE = re.compile("([\u2190-\u21ff\u2300-\u23ff\u2600-\u27bf\u2900-\u297f\u2b00-\u2bff"
                       "\u3030\u303d\u3297\u3299\U0001f000-\U0001faff\ufe0f\u200d\u20e3]+)")
_emoji_fonts = {}


def _emoji_font(size):
    """Font used for emoji; None if no emoji-capable font is available."""
    if size not in _emoji_fonts:
        found = None
        for cand in _EMOJI_FONTS:
            for s in ([size] if cand != "NotoColorEmoji.ttf" and "Noto" not in cand else [109]):
                try:
                    found = ImageFont.truetype(cand, s)
                    break
                except OSError:
                    continue
            if found:
                break
        _emoji_fonts[size] = found
    return _emoji_fonts[size]


def _segments(text, font):
    """Split text into (segment, font, is_emoji) so emoji use an emoji-capable font."""
    out = []
    for i, part in enumerate(_EMOJI_RE.split(text)):
        if not part:
            continue
        ef = _emoji_font(int(font.size)) if i % 2 else None
        out.append((part, ef, True) if ef else (part, font, False))
    return out


def _emoji_scale(ef, font):
    return font.size / ef.size if ef.size > font.size * 2 else 1.0


def _tlen(draw, text, font):
    """Text width, accounting for emoji fallback fonts."""
    total = 0.0
    for seg, f, emoji in _segments(text, font):
        w = draw.textlength(seg, font=f)
        total += w * _emoji_scale(f, font) if emoji else w
    return total


def _draw_run(img, draw, xy, text, font, fill):
    """Draw text, rendering emoji runs with an emoji font (scaled bitmap fonts included)."""
    x, y = xy
    for seg, f, emoji in _segments(text, font):
        if not emoji:
            draw.text((x, y), seg, font=f, fill=fill)
            x += draw.textlength(seg, font=f)
            continue
        scale = _emoji_scale(f, font)
        w = draw.textlength(seg, font=f)
        if scale == 1.0:
            try:
                draw.text((x, y), seg, font=f, fill=fill, embedded_color=True)
            except (OSError, ValueError):
                draw.text((x, y), seg, font=f, fill=fill)
        else:
            asc, desc = f.getmetrics()
            tile = Image.new("RGBA", (max(int(w) + 1, 1), asc + desc + 1), (0, 0, 0, 0))
            ImageDraw.Draw(tile).text((0, 0), seg, font=f, embedded_color=True)
            tile = tile.resize((max(round(tile.width * scale), 1), max(round(tile.height * scale), 1)),
                               Image.LANCZOS)
            img.paste(tile, (round(x), round(y)), tile)
        x += w * scale


_THEME_COLOR_KEYS = {
    "DARK_1": "dk1", "LIGHT_1": "lt1", "DARK_2": "dk2", "LIGHT_2": "lt2",
    "TEXT_1": "tx1", "BACKGROUND_1": "bg1", "TEXT_2": "tx2", "BACKGROUND_2": "bg2",
    "HYPERLINK": "hlink", "FOLLOWED_HYPERLINK": "folHlink",
    **{"ACCENT_%d" % i: "accent%d" % i for i in range(1, 7)},
}


def _rgb(color_format, theme_colors=None):
    """Return an RGB tuple for a python-pptx ColorFormat, or None."""
    try:
        if color_format.type is None:
            return None
        return tuple(color_format.rgb)
    except (AttributeError, KeyError, TypeError, ValueError):
        pass
    try:  # theme colors have no direct rgb
        theme_color = color_format.theme_color
        name = getattr(theme_color, "name", "")
        key = _THEME_COLOR_KEYS.get(name)
        if theme_colors and key in theme_colors:
            return theme_colors[key]
        return {
            "BACKGROUND_1": (255, 255, 255), "LIGHT_1": (255, 255, 255),
            "TEXT_1": (0, 0, 0), "DARK_1": (0, 0, 0),
        }.get(name)
    except Exception:
        return None


def _fill_color(fill, theme_colors=None):
    try:
        if fill.type == MSO_FILL.SOLID:
            return _rgb(fill.fore_color, theme_colors)
        if fill.type == MSO_FILL.GRADIENT:
            return _rgb(fill.gradient_stops[0].color, theme_colors)
    except Exception:
        pass
    return None


def _alpha(fill_parent, fill_tag="solidFill"):
    """Opacity (0-255) from an a:alpha child of the color in fill_parent/a:<fill_tag>; 255 if absent."""
    try:
        vals = fill_parent.xpath(f"./a:{fill_tag}/*/a:alpha/@val")
        if vals:
            return max(0, min(255, round(int(vals[0]) / 100000 * 255)))
    except Exception:
        pass
    return 255


def _with_alpha(color, alpha):
    return None if color is None else (*color[:3], alpha)


def _gradient_image(grad_fill, size, theme_colors):
    """Render a DrawingML linear gradient into an RGBA image."""
    stops_el = grad_fill.find(_A_NS + "gsLst")
    line = grad_fill.find(_A_NS + "lin")
    if stops_el is None or line is None:
        return None

    stops = []
    for stop in stops_el.findall(_A_NS + "gs"):
        if len(stop):
            color = _xml_color(stop[0], theme_colors)
            if color is not None:
                stops.append((int(stop.get("pos", "0")) / 100000, color))
    if len(stops) < 2:
        return None
    stops.sort(key=lambda stop: stop[0])

    width, height = size
    angle = math.radians(int(line.get("ang", "0")) / 60000)
    cos_a, sin_a = math.cos(angle), math.sin(angle)
    span = abs(cos_a) * max(width - 1, 0) + abs(sin_a) * max(height - 1, 0)
    bar_width = max(round(span) + 1, 1)

    def color_at(position):
        if position <= stops[0][0]:
            return stops[0][1]
        for (start, first), (end, second) in zip(stops, stops[1:]):
            if position <= end:
                ratio = (position - start) / (end - start) if end > start else 1
                return tuple(round(a + (b - a) * ratio) for a, b in zip(first, second))
        return stops[-1][1]

    bar = Image.new("RGBA", (bar_width, 1))
    bar.putdata([color_at(i / (bar_width - 1) if bar_width > 1 else 0)
                 for i in range(bar_width)])

    scale = (bar_width - 1) / span if span else 0
    x_scale, y_scale = cos_a * scale, sin_a * scale
    x_offset = (bar_width - 1) / 2 - x_scale * (width - 1) / 2 - y_scale * (height - 1) / 2
    return bar.transform(
        size, Image.Transform.AFFINE,
        (x_scale, y_scale, x_offset, 0, 0, 0),
        resample=Image.Resampling.BILINEAR,
    )


class _Canvas:
    def __init__(self, width_emu, height_emu, width_px):
        self.scale = width_px / width_emu
        self.size = (width_px, max(1, round(height_emu * self.scale)))
        self.image = Image.new("RGB", self.size, (255, 255, 255))
        self.prs = self.master = None
        self.theme_fonts = {}
        self.theme_colors = {}
        self.theme_lines = []
        self.table_styles = {}
        self.default_table_style = None

    def px(self, emu):
        return emu * self.scale

    def box(self, shape, ox=0, oy=0, sx=1.0, sy=1.0):
        """Pixel bounding box (x0, y0, x1, y1) of a shape, with group transform."""
        l, t, w, h = shape.left, shape.top, shape.width, shape.height
        if None in (l, t, w, h):
            return None
        x0 = ox + l * sx
        y0 = oy + t * sy
        return (self.px(x0), self.px(y0), self.px(x0 + w * sx), self.px(y0 + h * sy))

    def paste_rotated(self, layer, box, rotation):
        """Composite an RGBA layer (sized to box) onto the canvas, rotated about its center."""
        if rotation:
            layer = layer.rotate(-rotation, expand=True, resample=Image.BICUBIC)
        cx, cy = (box[0] + box[2]) / 2, (box[1] + box[3]) / 2
        pos = (round(cx - layer.width / 2), round(cy - layer.height / 2))
        self.image.paste(layer, pos, layer)


_ALIGN = {"l": PP_ALIGN.LEFT, "ctr": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT, "just": PP_ALIGN.JUSTIFY}
_TOKEN_RE = re.compile(r"\t|[\n\v]|[^ \t\n\v]+ *| +")
_DEFAULT_TAB_EMU = 914400
_A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _para_props(canvas, shape, para):
    """Resolve paragraph properties (inheriting from list styles, layout, master) as a dict:
    marL, indent (EMU), algn, defTabSz (EMU) and tabs ([(pos_emu, align)] sorted)."""
    chain = [para._p.pPr] if para._p.pPr is not None else []
    chain += list(_style_sources(canvas, shape, para.level, None))
    res = {"marL": 0, "indent": 0, "algn": None, "defTabSz": _DEFAULT_TAB_EMU, "tabs": []}
    seen = set()
    for el in chain:
        for key in ("marL", "indent", "algn", "defTabSz"):
            if key not in seen and el.get(key) is not None:
                seen.add(key)
                res[key] = el.get(key) if key == "algn" else int(el.get(key))
        if "tabs" not in seen:
            tabs = [(int(t.get("pos")), t.get("algn", "l")) for t in el.iter(_A_NS + "tab")
                    if t.get("pos") is not None]
            if tabs:
                seen.add("tabs")
                res["tabs"] = sorted(tabs)
    res["defTabSz"] = max(res["defTabSz"], 1)
    return res


_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_NS = {"a": _A}


def _theme_fonts(master):
    """Return {'+mj-lt': major latin font, '+mn-lt': minor latin font} of the master's theme."""
    try:
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT
        from lxml import etree
        theme = etree.fromstring(master.part.part_related_by(RT.THEME).blob)
        return {
            "+mj-lt": theme.find(".//a:majorFont/a:latin", _NS).get("typeface"),
            "+mn-lt": theme.find(".//a:minorFont/a:latin", _NS).get("typeface"),
        }
    except Exception:
        return {}


def _theme_colors(master):
    """Return the theme's named color scheme as RGB tuples."""
    try:
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT
        from lxml import etree
        theme = etree.fromstring(master.part.part_related_by(RT.THEME).blob)
        scheme = theme.find(".//a:clrScheme", _NS)
        colors = {}
        for entry in scheme:
            color = entry[0]
            value = color.get("lastClr") if color.tag == _A_NS + "sysClr" else color.get("val")
            if value:
                colors[entry.tag.rsplit("}", 1)[-1]] = tuple(
                    int(value[i:i + 2], 16) for i in (0, 2, 4))
        mapping = {"bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2"}
        clr_map = master._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}clrMap")
        if clr_map is not None:
            for key in mapping:
                if clr_map.get(key):
                    mapping[key] = clr_map.get(key)
        colors.update({key: colors.get(target) for key, target in mapping.items()})
        return colors
    except Exception:
        return {}


def _style_sources(canvas, shape, level, leaf="a:defRPr"):
    """Yield list-style elements (defRPr, or pPr itself if leaf is None) from most to least specific."""
    tag = "a:lvl%dpPr" % (level + 1) + ("/" + leaf if leaf else "")

    def lst(txbody):
        if txbody is not None:
            for el in txbody.xpath("a:lstStyle/" + tag):
                yield el

    kind = "other"
    if shape is not None and shape.is_placeholder:
        ph = shape.placeholder_format.type
        if ph in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            kind = "title"
        elif ph not in (PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER):
            kind = "body"
    cur = shape
    for _ in range(3):  # shape -> layout placeholder -> master placeholder
        if cur is None:
            break
        yield from lst(cur._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}txBody"))
        cur = getattr(cur, "_base_placeholder", None) if cur.is_placeholder else None
    if shape is not None and leaf == "a:defRPr":
        # the shape style's fontRef color outranks master/default text styles
        ref = shape._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style/a:fontRef", _NS)
        if ref is not None and len(ref):
            from lxml import etree
            rpr = etree.Element(_A_NS + "defRPr")
            etree.SubElement(rpr, _A_NS + "solidFill").append(copy.deepcopy(ref[0]))
            yield rpr
    master = canvas.master
    if master is not None:
        styles = master._element.xpath("p:txStyles/p:%sStyle" % kind)
        for st in styles:
            yield from st.xpath(tag, namespaces=_NS)
    if canvas.prs is not None:
        yield from canvas.prs._element.xpath("p:defaultTextStyle/" + tag)


def _vertical_anchor(text_frame, shape):
    anchor = text_frame.vertical_anchor
    current = shape
    while anchor is None and getattr(current, "is_placeholder", False):
        current = getattr(current, "_base_placeholder", None)
        if current is not None:
            anchor = current.text_frame.vertical_anchor
    return anchor


def _run_props(canvas, shape, para, run):
    """Resolve (size in EMU, bold, italic, font name) of a run, following inheritance."""
    chain = [run._r.rPr] if run._r.rPr is not None else []
    chain += list(_style_sources(canvas, shape, para.level))
    size = bold = italic = name = None
    for el in chain:
        if size is None and el.get("sz"):
            size = int(el.get("sz")) * 127  # hundredths of a point -> EMU
        if bold is None and el.get("b") is not None:
            bold = el.get("b") in ("1", "true")
        if italic is None and el.get("i") is not None:
            italic = el.get("i") in ("1", "true")
        if name is None:
            latin = el.find("a:latin", _NS)
            if latin is not None and latin.get("typeface"):
                name = latin.get("typeface")
    name = canvas.theme_fonts.get(name, name) if name else canvas.theme_fonts.get("+mn-lt")
    return (size or 18 * 12700), bool(bold), bool(italic), name


def _font_scale(text_frame):
    """Scale applied by PowerPoint's 'shrink text on overflow' (normAutofit fontScale)."""
    try:
        el = text_frame._txBody.bodyPr.find("a:normAutofit", _NS)
        if el is not None and el.get("fontScale"):
            return int(el.get("fontScale")) / 100000.0
    except Exception:
        pass
    return 1.0


def _text_rotation(shape_rotation, text_frame):
    """Total rotation (degrees) of text: shape rotation plus bodyPr rot."""
    rot = shape_rotation or 0
    try:
        body_rot = text_frame._txBody.bodyPr.get("rot")
        if body_rot:
            rot += int(body_rot) / 60000.0
    except Exception:
        pass
    return rot % 360


def _run_color(canvas, shape, para, run, default_color):
    """Resolve a run's explicit or inherited solid text color."""
    chain = [run._r.rPr] if run._r.rPr is not None else []
    chain += list(_style_sources(canvas, shape, para.level))
    for el in chain:
        solid_fill = el.find("a:solidFill", _NS)
        if solid_fill is None or len(solid_fill) == 0:
            continue
        color = _xml_color(solid_fill[0], canvas.theme_colors)
        if color is not None:
            return color[:3]
    return default_color


def _layout_line(draw, tokens, canvas, stops, default_tab, start=0.0):
    """Place tokens of one line; returns [(x, text, font, color)] and the line width."""
    placed, x = [], start
    i = 0
    while i < len(tokens):
        text, font, color, _, ul = tokens[i]
        if text == "\t":
            nxt = next(((canvas.px(p), al) for p, al in stops if canvas.px(p) > x + 0.5), None)
            if nxt is None:
                step = canvas.px(default_tab)
                nxt = ((x // step + 1) * step, "l")
            pos, align = nxt
            if align in ("ctr", "r", "dec"):
                seg = 0.0
                for t, f, _c, _s, _u in tokens[i + 1:]:
                    if t == "\t":
                        break
                    seg += _tlen(draw, t, f)
                x = max(pos - (seg / 2 if align == "ctr" else seg), x)
            else:
                x = pos
        else:
            placed.append((x, text, font, color, ul))
            x += _tlen(draw, text, font)
        i += 1
    return placed, x


class _RunShim:
    """Uniform view of a:r, a:br and a:fld paragraph children."""

    def __init__(self, el):
        self._r = self
        self.el = el
        self.tag = el.tag
        self.rPr = el.find("a:rPr", _NS)
        t = el.find("a:t", _NS)
        self.text = (t.text or "") if t is not None else ""
        self.field_type = el.get("type") if el.tag == _A_NS + "fld" else None


def _para_runs(para):
    for el in para._p:
        if el.tag in (_A_NS + "r", _A_NS + "br", _A_NS + "fld"):
            yield _RunShim(el)


def _xml_color(el, theme_colors):
    """RGBA of a color element (srgbClr/schemeClr/...) with lumMod/lumOff/tint/shade/alpha applied."""
    import colorsys
    value = el.get("lastClr") if el.tag == _A_NS + "sysClr" else el.get("val")
    if el.tag == _A_NS + "schemeClr":
        rgb = theme_colors.get(value)
    elif el.tag in (_A_NS + "srgbClr", _A_NS + "sysClr") and value:
        rgb = tuple(int(value[i:i + 2], 16) for i in (0, 2, 4))
    else:
        rgb = None
    if rgb is None:
        return None
    alpha = 255
    rgb = tuple(c / 255 for c in rgb)
    for mod in el:
        v = int(mod.get("val", "0")) / 100000
        name = mod.tag.rsplit("}", 1)[-1]
        if name == "lumMod":
            h, l, s = colorsys.rgb_to_hls(*rgb)
            l = min(l * v, 1)
            rgb = colorsys.hls_to_rgb(h, l, s)
        elif name == "lumOff":
            h, l, s = colorsys.rgb_to_hls(*rgb)
            l = min(max(l + v, 0), 1)
            rgb = colorsys.hls_to_rgb(h, l, s)
        elif name == "alpha":
            alpha = round(min(max(v, 0), 1) * 255)
        elif name == "tint":
            rgb = tuple(c + (1 - c) * v for c in rgb)
        elif name == "shade":
            rgb = tuple(c * v for c in rgb)
        elif name == "satMod":
            h, l, s = colorsys.rgb_to_hls(*rgb)
            rgb = colorsys.hls_to_rgb(h, l, min(max(s * v, 0), 1))
    return (*tuple(round(c * 255) for c in rgb), alpha)


def _shape_fill(shape, theme_colors):
    """RGBA fill of a shape: explicit spPr fill, else the style's fillRef; None if no fill."""
    try:
        sppr = shape._element.spPr
    except AttributeError:
        return None
    if sppr.find("a:noFill", _NS) is not None:
        return None
    for tag in ("a:solidFill", "a:gradFill"):
        fill = sppr.find(tag, _NS)
        if fill is None:
            continue
        if tag == "a:gradFill":
            return None
        if fill is not None and len(fill):
            return _xml_color(fill[0], theme_colors)
        return None
    if sppr.find("a:blipFill", _NS) is not None or sppr.find("a:pattFill", _NS) is not None:
        return None
    ref = shape._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style/a:fillRef", _NS)
    if ref is not None and ref.get("idx", "0") != "0" and len(ref):
        return _xml_color(ref[0], theme_colors)
    return None


def _shape_gradient(shape):
    try:
        return shape._element.spPr.find(_A_NS + "gradFill")
    except AttributeError:
        return None


def _draw_text(canvas, text_frame, box, default_color=(0, 0, 0), rotation=0, shape=None):
    if rotation:
        # draw on a padded transparent layer centred on the box, then rotate it onto the canvas
        w, h = box[2] - box[0], box[3] - box[1]
        pad = int(max(w, h))
        layer = Image.new("RGBA", (int(w) + 2 * pad + 1, int(h) + 2 * pad + 1), (0, 0, 0, 0))
        local = (pad, pad, pad + w, pad + h)
        _render_text(canvas, layer, text_frame, local, default_color, shape)
        canvas.paste_rotated(layer, box, rotation)
    else:
        _render_text(canvas, canvas.image, text_frame, box, default_color, shape)


def _render_text(canvas, target, text_frame, box, default_color, shape=None):
    x0, y0, x1, y1 = box
    pad_l = canvas.px(text_frame.margin_left)
    pad_r = canvas.px(text_frame.margin_right)
    pad_t = canvas.px(text_frame.margin_top)
    pad_b = canvas.px(text_frame.margin_bottom)
    avail = max(x1 - x0 - pad_l - pad_r, 1)
    draw = ImageDraw.Draw(target)

    lines = []  # (placed items, line width, align, height)
    for para in text_frame.paragraphs:
        props = _para_props(canvas, shape, para)
        stops, default_tab = props["tabs"], props["defTabSz"]
        mar_l, indent = props["marL"], props["indent"]
        align = _ALIGN.get(props["algn"], para.alignment)
        if mar_l and indent < 0:  # hanging indent: the left margin acts as a tab stop
            stops = sorted(stops + [(mar_l, "l")])
        first_start = canvas.px(max(mar_l + indent, 0))
        rest_start = canvas.px(mar_l)
        line_no = 0
        tokens = []
        for r in _para_runs(para):
            is_br = r._r.tag == _A_NS + "br"
            text = r.text
            if r.field_type == "slidenum" and getattr(canvas, "slide_number", None):
                text = str(canvas.slide_number)
            if not is_br and not text:
                continue
            size, bold, italic, name = _run_props(canvas, shape, para, r)
            # always render 1pt smaller than specified
            size = max(size * _font_scale(text_frame) - 12700, 12700)
            fpx = canvas.px(size)
            font = _font(fpx, bold, name, italic)
            color = _run_color(canvas, shape, para, r, default_color)
            rpr = r._r.rPr
            underline = rpr is not None and rpr.get("u", "none") != "none"
            if rpr is not None and rpr.find("a:hlinkClick", _NS) is not None:
                color = canvas.theme_colors.get("hlink") or (5, 99, 193)
                underline = True
            if is_br:
                tokens.append(("\n", font, color, fpx, False))
                continue
            for tok in _TOKEN_RE.findall(text):
                tokens.append((tok, font, color, fpx, underline))
        if not tokens:
            h = canvas.px(Emu(127000))
            lines.append(([], 0, align, h * 1.2))
            continue

        cur = []

        def flush():
            nonlocal cur, line_no
            if cur:
                start = first_start if line_no == 0 else rest_start
                line_no += 1
                placed, width = _layout_line(draw, cur, canvas, stops, default_tab, start)
                width = max(
                    (x + _tlen(draw, t.rstrip(" "), f) for x, t, f, _c, _u in placed), default=0)
                lines.append((placed, width, align, max(t[3] for t in cur) * 1.2))
            cur = []

        for tok in tokens:
            text, font = tok[0], tok[1]
            if text in ("\n", "\v"):
                if cur:
                    flush()
                else:  # consecutive break: empty line
                    lines.append(([], 0, align, tok[3] * 1.2))
                continue
            if text != "\t" and cur:
                start = first_start if line_no == 0 else rest_start
                _p, cur_w = _layout_line(draw, cur, canvas, stops, default_tab, start)
                if cur_w + _tlen(draw, text.rstrip(" "), font) > avail:
                    flush()
            cur.append(tok)
        flush()

    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    total = sum(l[3] for l in lines)
    anchor = _vertical_anchor(text_frame, shape)
    inner_h = y1 - y0 - pad_t - pad_b
    y = y0 + pad_t
    if anchor == MSO_ANCHOR.MIDDLE:
        y += (inner_h - total) / 2
    elif anchor == MSO_ANCHOR.BOTTOM:
        y += inner_h - total
    for placed, line_w, align, h in lines:
        shift = 0
        if align == PP_ALIGN.CENTER:
            shift = (avail - line_w) / 2
        elif align == PP_ALIGN.RIGHT:
            shift = avail - line_w
        for x, text, font, color, ul in placed:
            tx, shown = x0 + pad_l + shift + x, text.rstrip(" ")
            _draw_run(target, draw, (tx, y), shown, font, color)
            if ul and shown:
                asc, _desc = font.getmetrics()
                uy = y + asc + max(font.size * 0.08, 1)
                draw.line([(tx, uy), (tx + _tlen(draw, shown, font), uy)],
                          fill=color, width=max(round(font.size / 14), 1))
        y += h


def _load_metafile(img, target_px=1600):
    """Rasterize an EMF/WMF at high resolution (Windows only); white becomes transparent."""
    try:
        w, h = img.size
        dpi = img.info.get("dpi", 72)
        dpi = dpi[0] if isinstance(dpi, tuple) else dpi
        img.load(dpi=dpi * max(target_px / max(w, h, 1), 1))
    except Exception:
        return None
    img = img.convert("RGBA")
    # the metafile is drawn onto an opaque white background
    r, g, b, _ = img.split()
    is_white = [c.point(lambda v: 255 if v == 255 else 0) for c in (r, g, b)]
    white = ImageChops.multiply(ImageChops.multiply(is_white[0], is_white[1]), is_white[2])
    img.putalpha(ImageChops.invert(white))
    return img


def _picture_image(shape):
    """Return the picture cropped according to its crop information."""
    img = Image.open(io.BytesIO(shape.image.blob))
    if img.format in ("WMF", "EMF"):
        img = _load_metafile(img)
        if img is None:
            return None
    else:
        img.load()
        img = img.convert("RGBA")
    w, h = img.size
    l, t, r, b = shape.crop_left, shape.crop_top, shape.crop_right, shape.crop_bottom
    # negative crop values pad the image with transparency
    x0, y0, x1, y1 = round(l * w), round(t * h), round(w - r * w), round(h - b * h)
    if x1 <= x0 or y1 <= y0:
        return None
    if min(x0, y0) >= 0 and x1 <= w and y1 <= h:
        return img.crop((x0, y0, x1, y1))
    out = Image.new("RGBA", (x1 - x0, y1 - y0), (0, 0, 0, 0))
    out.paste(img, (-x0, -y0))
    return out


def _apply_duotone(img, shape, theme_colors):
    """Recolor a picture per a:duotone: luminance 0 maps to the first color, 1 to the second."""
    try:
        duotone = shape._element.blipFill.find("a:blip/a:duotone", _NS)
    except AttributeError:
        return img
    if duotone is None or len(duotone) < 2:
        return img
    c0 = _xml_color(duotone[0], theme_colors) if duotone[0].tag != _A_NS + "prstClr" else None
    c1 = _xml_color(duotone[1], theme_colors) if duotone[1].tag != _A_NS + "prstClr" else None
    preset = {"white": (255, 255, 255, 255), "black": (0, 0, 0, 255)}
    if c0 is None:
        c0 = preset.get(duotone[0].get("val"))
    if c1 is None:
        c1 = preset.get(duotone[1].get("val"))
    if c0 is None or c1 is None:
        return img
    alpha = img.getchannel("A")
    gray = img.convert("L")
    channels = [gray.point([round(c0[i] + (c1[i] - c0[i]) * v / 255) for v in range(256)])
                for i in range(3)]
    out = Image.merge("RGB", channels).convert("RGBA")
    out.putalpha(alpha)
    return out


def _draw_picture(canvas, shape, box):
    img = _picture_image(shape)
    if img is None:
        return
    img = _apply_duotone(img, shape, canvas.theme_colors)
    bw, bh = max(round(box[2] - box[0]), 1), max(round(box[3] - box[1]), 1)
    img = img.resize((bw, bh), Image.LANCZOS)
    canvas.paste_rotated(img, box, getattr(shape, "rotation", 0))


_SHADE = {"darken": 0.6, "darkenLess": 0.8, "lighten": -0.4, "lightenLess": -0.2}


def _shade(color, mode):
    f = _SHADE.get(mode)
    if color is None or f is None:
        return color
    rgb = [round(c * f) if f > 0 else round(c + (255 - c) * -f) for c in color[:3]]
    return tuple(rgb) + tuple(color[3:])


def _theme_lines(master):
    """Return the theme's line styles (a:ln elements of fmtScheme/lnStyleLst)."""
    try:
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT
        from lxml import etree
        theme = etree.fromstring(master.part.part_related_by(RT.THEME).blob)
        return list(theme.find(".//a:fmtScheme/a:lnStyleLst", _NS))
    except Exception:
        return []


_P_STYLE = "{http://schemas.openxmlformats.org/presentationml/2006/main}style"
_JOINS = {"round": "round", "bevel": "bevel", "miter": "miter"}


def _shape_line(canvas, shape):
    """Resolved outline of a shape: dict(color, width, dash, cap, join, miter, head, tail) or None."""
    from .lines import DASHES
    el = shape._element
    sppr = getattr(el, "spPr", None)
    if sppr is None:
        return None
    ln = sppr.find(_A_NS + "ln")
    style = el.find(_P_STYLE)
    ref = style.find(_A_NS + "lnRef") if style is not None else None
    idx = int(ref.get("idx", "0")) if ref is not None else 0
    theme_ln = canvas.theme_lines[idx - 1] if 1 <= idx <= len(canvas.theme_lines) else None
    sources = [s for s in (ln, theme_ln) if s is not None]
    if ln is not None and ln.find(_A_NS + "noFill") is not None:
        return None

    color = None
    fill_el = None
    if ln is not None:
        for tag in ("solidFill", "gradFill", "pattFill"):
            fill_el = ln.find(_A_NS + tag)
            if fill_el is not None:
                fill_tag = tag
                break
    if fill_el is not None:
        if fill_tag == "gradFill":
            fill_el = fill_el.find(_A_NS + "gsLst/" + _A_NS + "gs")
        elif fill_tag == "pattFill":
            fill_el = fill_el.find(_A_NS + "fgClr")
        if fill_el is not None and len(fill_el):
            color = _xml_color(fill_el[0], canvas.theme_colors)
    elif ref is not None and idx > 0 and len(ref):
        color = _xml_color(ref[0], canvas.theme_colors)
    if color is None:
        return None

    def attr(name):
        for s in sources:
            if s.get(name) is not None:
                return s.get(name)
        return None

    def child(tag):
        for s in sources:
            c = s.find(_A_NS + tag)
            if c is not None:
                return c
        return None

    width = max(canvas.px(int(attr("w") or 12700)), 1.0)
    dash = None
    pd = child("prstDash")
    cd = child("custDash")
    if pd is not None:
        dash = DASHES.get(pd.get("val"))
    elif cd is not None:
        vals = []
        for ds in cd:
            vals += [int(ds.get("d", "100000")) / 100000.0, int(ds.get("sp", "100000")) / 100000.0]
        dash = tuple(vals) or None
    join, miter = "round", 8.0
    for s in sources:
        found = [c for c in s if c.tag.rsplit("}", 1)[-1] in _JOINS]
        if found:
            join = _JOINS[found[0].tag.rsplit("}", 1)[-1]]
            miter = int(found[0].get("lim", "800000")) / 100000.0
            break
    head = tail = None
    for tag in ("headEnd", "tailEnd"):
        end = ln.find(_A_NS + tag) if ln is not None else None
        if end is not None and end.get("type", "none") != "none":
            info = (end.get("type"), {"sm": 2, "med": 3, "lg": 5}.get(end.get("w", "med"), 3))
            if tag == "headEnd":
                head = info
            else:
                tail = info
    return {"color": color, "width": width, "dash": dash,
            "cap": {"rnd": "rnd", "sq": "sq"}.get(attr("cap"), "flat"),
            "join": join, "miter": miter, "head": head, "tail": tail}


def _arrow_ref(pts, i, dist):
    """First point at least ``dist`` away from pts[i] (to get a stable arrow direction)."""
    base = pts[i]
    for p in pts[i + 1:]:
        if math.hypot(p[0] - base[0], p[1] - base[1]) >= dist:
            return p
    return pts[-1]


def _trim_polyline(pts, dist):
    """Drop ``dist`` of length from the start of a polyline."""
    left = dist
    for i in range(len(pts) - 1):
        seg = math.hypot(pts[i + 1][0] - pts[i][0], pts[i + 1][1] - pts[i][1])
        if seg > left:
            t = left / seg if seg else 0
            start = (pts[i][0] + (pts[i + 1][0] - pts[i][0]) * t, pts[i][1] + (pts[i + 1][1] - pts[i][1]) * t)
            return [start] + pts[i + 1:]
        left -= seg
    return pts


def _trim_for_arrows(pts, width, head, tail):
    """Shorten the stroked body at arrow ends so its thickness doesn't poke out past the tip."""
    cut = width * 1.5
    if head:
        pts = _trim_polyline(pts, cut)
    if tail:
        pts = _trim_polyline(pts[::-1], cut)[::-1]
    return pts


def _draw_shape_geometry(canvas, shape, box, fill, line, gradient_fill=None, theme_colors=None):
    from . import geometry, lines
    ss = 2
    bw, bh = max(round(box[2] - box[0]), 1), max(round(box[3] - box[1]), 1)
    line_w = line["width"] if line else 0
    inset = 0.0
    w, h = bw * ss, bh * ss
    geom, overrides = geometry.shape_geometry(shape)
    paths = geometry.evaluate(geom, w, h, overrides) if geom is not None else None
    if not paths:
        paths = [{"fill": "norm", "stroke": True, "subpaths": [([(0, 0), (w, 0), (w, h), (0, h)], True)]}]
    xfrm = shape._element.spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm")
    flip_h = xfrm is not None and xfrm.get("flipH") in ("1", "true")
    flip_v = xfrm is not None and xfrm.get("flipV") in ("1", "true")

    def place(pts):
        return [((w - x if flip_h else x) + inset, (h - y if flip_v else y) + inset) for x, y in pts]

    placed = [[(place(pts), closed) for pts, closed in path["subpaths"]] for path in paths]
    # callout tails etc. may extend beyond the shape box: pad symmetrically so rotation stays centred
    xs = [x for subs in placed for pts, _ in subs for x, _ in pts] or [0]
    ys = [y for subs in placed for pts, _ in subs for _, y in pts] or [0]
    reach = line_w * ss * (min(line["miter"], 6.0) if line and line["join"] == "miter" else 1.0)
    over_x = max(0.0, -min(xs), max(xs) - bw * ss) + reach
    over_y = max(0.0, -min(ys), max(ys) - bh * ss) + reach
    pad_x, pad_y = math.ceil(over_x / ss), math.ceil(over_y / ss)
    layer = Image.new("RGBA", ((bw + 2 * pad_x) * ss, (bh + 2 * pad_y) * ss), (0, 0, 0, 0))
    off = (pad_x * ss, pad_y * ss)
    box = (box[0] - pad_x, box[1] - pad_y, box[2] + pad_x, box[3] + pad_y)
    gradient = None
    if gradient_fill is not None:
        gradient = _gradient_image(gradient_fill, (bw * ss, bh * ss), theme_colors or {})
        if gradient is not None:
            gradient_layer = Image.new("RGBA", layer.size, (0, 0, 0, 0))
            gradient_layer.paste(gradient, off)
            gradient = gradient_layer
        else:
            stops = gradient_fill.find(_A_NS + "gsLst")
            first = stops.find(_A_NS + "gs") if stops is not None else None
            if first is not None and len(first):
                fill = _xml_color(first[0], theme_colors or {})

    for path, subs in zip(paths, placed):
        subs = [([(x + off[0], y + off[1]) for x, y in pts], closed) for pts, closed in subs]
        if (fill or gradient) and path["fill"] != "none":
            # even-odd fill so that nested sub-paths (e.g. donut) become holes
            mask = Image.new("1", layer.size, 0)
            for pts, _ in subs:
                if len(pts) >= 3:
                    tmp = Image.new("1", layer.size, 0)
                    ImageDraw.Draw(tmp).polygon(pts, fill=1)
                    mask = ImageChops.logical_xor(mask, tmp)
            layer.paste(gradient if gradient else _shade(fill, path["fill"]), mask=mask)
        if line and path["stroke"]:
            for pts, closed in subs:
                if len(pts) < 2:
                    continue
                head = None if closed else line["head"]
                tail = None if closed else line["tail"]
                body = _trim_for_arrows(pts, line_w * ss, head, tail)
                lines.stroke(layer, body, closed, line["color"], line_w * ss, line["dash"],
                             line["cap"], line["join"], line["miter"])
                d = ImageDraw.Draw(layer)
                if head:
                    _draw_arrow_head(d, pts[0], _arrow_ref(pts, 0, line_w * ss), head[0], head[1],
                                     line_w * ss, line["color"])
                if tail:
                    _draw_arrow_head(d, pts[-1], _arrow_ref(pts[::-1], 0, line_w * ss), tail[0], tail[1],
                                     line_w * ss, line["color"])
    layer = layer.resize((bw + 2 * pad_x, bh + 2 * pad_y), Image.LANCZOS)
    canvas.paste_rotated(layer, box, getattr(shape, "rotation", 0))


def _table_styles(prs):
    """Return table styles in the presentation and its default table-style ID."""
    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    try:
        part = prs.part.part_related_by(RT.TABLE_STYLES)
    except KeyError:
        return {}, None
    root = etree.fromstring(part.blob)
    styles = {
        style.get("styleId").strip("{}").lower(): style
        for style in root.findall("a:tblStyle", _NS)
        if style.get("styleId")
    }
    default = root.get("def")
    return styles, default.strip("{}").lower() if default else None


def _table_style_for_id(style_id, styles, default_style_id=None):
    """Use a defined table style, or the presentation's declared default."""
    style = styles.get(style_id)
    return style if style is not None else styles.get(default_style_id)


def _table_style_regions(tbl_pr, row, col, row_count, col_count):
    regions = []
    if row == 0 and tbl_pr.get("firstRow") in ("1", "true"):
        regions.append("firstRow")
    if row == row_count - 1 and tbl_pr.get("lastRow") in ("1", "true"):
        regions.append("lastRow")
    if col == 0 and tbl_pr.get("firstCol") in ("1", "true"):
        regions.append("firstCol")
    if col == col_count - 1 and tbl_pr.get("lastCol") in ("1", "true"):
        regions.append("lastCol")
    if tbl_pr.get("bandRow") in ("1", "true"):
        band_row = row - (1 if tbl_pr.get("firstRow") in ("1", "true") else 0)
        if band_row >= 0:
            regions.append("band%dH" % (band_row % 2 + 1))
    if tbl_pr.get("bandCol") in ("1", "true"):
        band_col = col - (1 if tbl_pr.get("firstCol") in ("1", "true") else 0)
        if band_col >= 0:
            regions.append("band%dV" % (band_col % 2 + 1))
    regions.append("wholeTbl")
    return regions


def _table_style_fill(style, tbl_pr, row, col, row_count, col_count, theme_colors):
    """Return the base RGB color declared for the applicable table-style fill."""
    if style is None:
        return None
    for name in _table_style_regions(tbl_pr, row, col, row_count, col_count):
        region = style.find("a:" + name + "/a:tcStyle", _NS)
        fill = region.find("a:fill", _NS) if region is not None else None
        color = fill.find("a:solidFill/*", _NS) if fill is not None else None
        if color is None:
            continue
        value = color.get("lastClr") if color.tag == _A_NS + "sysClr" else color.get("val")
        if color.tag == _A_NS + "schemeClr":
            return theme_colors.get(value)
        if color.tag in (_A_NS + "srgbClr", _A_NS + "sysClr") and value:
            return tuple(int(value[i:i + 2], 16) for i in (0, 2, 4))
    return None


def _draw_table(canvas, shape, box):
    table = shape.table
    tbl_pr = table._tbl.tblPr
    style_id = tbl_pr.find("a:tableStyleId", _NS)
    style_id = style_id.text.strip("{}").lower() if style_id is not None and style_id.text else None
    style = _table_style_for_id(
        style_id, canvas.table_styles, canvas.default_table_style)
    draw = ImageDraw.Draw(canvas.image)
    y = box[1]
    for ri, row in enumerate(table.rows):
        rh = canvas.px(row.height)
        x = box[0]
        for ci, col in enumerate(table.columns):
            cw = canvas.px(col.width)
            cell = table.cell(ri, ci)
            if cell.is_spanned:
                x += cw
                continue
            tcpr = cell._tc.tcPr
            if tcpr is not None and tcpr.find("a:noFill", _NS) is not None:
                c = None
            else:
                gradient_drawn = False
                solid_fill = tcpr.find("a:solidFill", _NS) if tcpr is not None else None
                if solid_fill is not None and len(solid_fill):
                    c = _xml_color(solid_fill[0], canvas.theme_colors)
                else:
                    grad_fill = tcpr.find("a:gradFill", _NS) if tcpr is not None else None
                    gradient = (_gradient_image(
                        grad_fill, (max(round(cw), 1), max(round(rh), 1)), canvas.theme_colors)
                        if grad_fill is not None else None)
                    if gradient is not None:
                        canvas.image.paste(gradient, (round(x), round(y)), gradient)
                        gradient_drawn = True
                        c = None
                    else:
                        c = _fill_color(cell.fill, canvas.theme_colors)
                    if c is None and not gradient_drawn:
                        c = _table_style_fill(
                            style, tbl_pr, ri, ci, len(table.rows), len(table.columns),
                            canvas.theme_colors)
            draw.rectangle((x, y, x + cw, y + rh), fill=c, outline=(128, 128, 128))
            _draw_text(canvas, cell.text_frame, (x, y, x + cw, y + rh))
            x += cw
        y += rh


def _draw_shapes(canvas, shapes, ox=0, oy=0, sx=1.0, sy=1.0, skip_placeholders=False):
    for shape in shapes:
        if skip_placeholders and shape.is_placeholder:
            continue
        try:
            _draw_shape(canvas, shape, ox, oy, sx, sy)
        except Exception:  # never fail a whole slide for one odd shape
            continue


_A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _draw_arrow_head(draw, tip, frm, kind, size_w, width, fill):
    import math
    dx, dy = tip[0] - frm[0], tip[1] - frm[1]
    length = math.hypot(dx, dy)
    if length == 0:
        return
    ux, uy = dx / length, dy / length
    nx, ny = -uy, ux
    s = max(size_w * width, 6)
    bx, by = tip[0] - ux * s, tip[1] - uy * s
    if kind == "oval":
        r = s / 2
        cx, cy = tip[0] - ux * r, tip[1] - uy * r
        draw.ellipse((cx - r, cy - r, cx + r, cy + r), fill=fill)
    elif kind == "diamond":
        r = s / 2
        cx, cy = tip[0] - ux * r, tip[1] - uy * r
        draw.polygon([tip, (cx + nx * r, cy + ny * r), (bx, by), (cx - nx * r, cy - ny * r)], fill=fill)
    elif kind in ("stealth", "arrow"):
        notch = s * (0.7 if kind == "stealth" else 1.0)
        nxp, nyp = tip[0] - ux * notch, tip[1] - uy * notch
        if kind == "stealth":
            draw.polygon([tip, (bx + nx * s / 2, by + ny * s / 2), (nxp, nyp),
                          (bx - nx * s / 2, by - ny * s / 2)], fill=fill)
        else:
            w = max(width, 2)
            for sgn in (1, -1):
                draw.line((tip[0], tip[1], bx + sgn * nx * s / 2, by + sgn * ny * s / 2), fill=fill, width=w)
    else:  # triangle
        draw.polygon([tip, (bx + nx * s / 2, by + ny * s / 2), (bx - nx * s / 2, by - ny * s / 2)], fill=fill)


def _draw_line_shape(canvas, shape, box):
    import math
    from . import lines
    info = _shape_line(canvas, shape) or {"color": (0, 0, 0, 255), "width": 1.0, "dash": None, "cap": "flat",
                                          "head": None, "tail": None}
    color = info["color"]
    width = max(round(info["width"]), 1)
    xfrm = shape._element.spPr.find(_A_NS + "xfrm")
    flip_h = flip_v = False
    rot = 0.0
    if xfrm is not None:
        flip_h = xfrm.get("flipH") in ("1", "true")
        flip_v = xfrm.get("flipV") in ("1", "true")
        rot = int(xfrm.get("rot", 0)) / 60000.0
    x0, y0, x1, y1 = box
    p0 = [x1 if flip_h else x0, y1 if flip_v else y0]
    p1 = [x0 if flip_h else x1, y0 if flip_v else y1]
    if rot:
        cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
        c, s = math.cos(math.radians(rot)), math.sin(math.radians(rot))
        for p in (p0, p1):
            rx, ry = p[0] - cx, p[1] - cy
            p[0], p[1] = cx + rx * c - ry * s, cy + rx * s + ry * c
    dash, head, tail = info["dash"], info["head"], info["tail"]
    a, b = tuple(p0), tuple(p1)
    length = math.hypot(b[0] - a[0], b[1] - a[1])
    # shorten the line at arrow ends so its thickness doesn't poke out past the tip
    la, lb = a, b
    if length:
        ux, uy = (b[0] - a[0]) / length, (b[1] - a[1]) / length
        cut_a = width * 1.5 if head else 0
        cut_b = width * 1.5 if tail else 0
        if cut_a + cut_b < length:
            la = (a[0] + ux * cut_a, a[1] + uy * cut_a)
            lb = (b[0] - ux * cut_b, b[1] - uy * cut_b)
    pad = int(width * 2 + 2)
    lx0, ly0 = int(min(a[0], b[0])) - pad, int(min(a[1], b[1])) - pad
    layer = Image.new("RGBA", (int(abs(b[0] - a[0])) + 2 * pad + 2, int(abs(b[1] - a[1])) + 2 * pad + 2), (0, 0, 0, 0))
    ss = 2
    big = layer.resize((layer.width * ss, layer.height * ss))
    loc = lambda p: ((p[0] - lx0) * ss, (p[1] - ly0) * ss)
    lines.stroke(big, [loc(la), loc(lb)], False, color, info["width"] * ss if info["width"] >= 1 else ss,
                 dash, info["cap"])
    if length:
        d = ImageDraw.Draw(big)
        if head:
            _draw_arrow_head(d, loc(a), loc(b), head[0], head[1], width * ss, color)
        if tail:
            _draw_arrow_head(d, loc(b), loc(a), tail[0], tail[1], width * ss, color)
    layer = big.resize(layer.size, Image.LANCZOS)
    canvas.image.paste(layer, (lx0, ly0), layer)


def _draw_shape(canvas, shape, ox, oy, sx, sy):
    st = shape.shape_type
    if st == MSO_SHAPE_TYPE.GROUP:
        # map child coordinate space onto the group's extents
        xfrm = shape._element.grpSpPr.xfrm
        ch_off, ch_ext = xfrm.chOff, xfrm.chExt
        gx = ox + shape.left * sx
        gy = oy + shape.top * sy
        nsx = shape.width * sx / ch_ext.cx if ch_ext.cx else sx
        nsy = shape.height * sy / ch_ext.cy if ch_ext.cy else sy
        _draw_shapes(canvas, shape.shapes, gx - ch_off.x * nsx, gy - ch_off.y * nsy, nsx, nsy)
        return
    box = canvas.box(shape, ox, oy, sx, sy)
    if box is None:
        return
    if st == MSO_SHAPE_TYPE.PICTURE or (shape.is_placeholder and hasattr(shape, "image")):
        _draw_picture(canvas, shape, box)
        return
    if getattr(shape, "has_table", False) and shape.has_table:
        _draw_table(canvas, shape, box)
        return
    if st == MSO_SHAPE_TYPE.LINE or shape.__class__.__name__ == "Connector":
        from . import geometry
        geom_el = geometry.shape_geometry(shape)[0]
        prst = None
        try:
            prst = shape._element.spPr.find(_A_NS + "prstGeom").get("prst")
        except AttributeError:
            pass
        if geom_el is None or prst in (None, "line", "straightConnector1"):
            _draw_line_shape(canvas, shape, box)
            return
        line = _shape_line(canvas, shape)
        if line:
            _draw_shape_geometry(canvas, shape, box, None, line)
        return
    fill = gradient_fill = None
    try:
        sppr = shape._element.spPr
        gradient_fill = _shape_gradient(shape)
        fill = _shape_fill(shape, canvas.theme_colors)
        if gradient_fill is None and fill is None and shape.fill.type is not None:
            fill = _with_alpha(_fill_color(shape.fill, canvas.theme_colors), _alpha(sppr))
    except Exception:
        pass
    line = _shape_line(canvas, shape)
    if fill or gradient_fill is not None or line:
        _draw_shape_geometry(canvas, shape, box, fill, line, gradient_fill, canvas.theme_colors)
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        _draw_text(canvas, shape.text_frame, box,
                   rotation=_text_rotation(getattr(shape, "rotation", 0), shape.text_frame), shape=shape)


def _background(canvas, slide):
    for obj in (slide, slide.slide_layout, slide.slide_layout.slide_master):
        try:
            bg_el = obj._element.cSld.bg
            bg_ref = bg_el.find("{http://schemas.openxmlformats.org/presentationml/2006/main}bgRef") \
                if bg_el is not None else None
            if bg_ref is not None and len(bg_ref):
                color = _xml_color(bg_ref[0], canvas.theme_colors)
                if color is not None:
                    canvas.image.paste(color[:3], (0, 0, *canvas.size))
                    return
            fill = obj.background.fill
            bg = obj._element.cSld.bg
            bg_pr = bg.bgPr if bg is not None else None
            grad_fill = bg_pr.find(_A_NS + "gradFill") if bg_pr is not None else None
            if grad_fill is not None:
                gradient = _gradient_image(grad_fill, canvas.size, canvas.theme_colors)
                if gradient is not None:
                    canvas.image = Image.alpha_composite(
                        canvas.image.convert("RGBA"), gradient
                    ).convert("RGB")
                    return
            c = _fill_color(fill, canvas.theme_colors)
            if c:
                canvas.image.paste(c, (0, 0, *canvas.size))
                return
        except Exception:
            continue


def _render(prs, slide, width):
    canvas = _Canvas(prs.slide_width, prs.slide_height, width)
    canvas.prs = prs
    canvas.master = slide.slide_layout.slide_master
    canvas.slide_number = next((i for i, s in enumerate(prs.slides, 1) if s.slide_id == slide.slide_id), None)
    canvas.theme_fonts = _theme_fonts(canvas.master)
    canvas.theme_colors = _theme_colors(canvas.master)
    canvas.theme_lines = _theme_lines(canvas.master)
    canvas.table_styles, canvas.default_table_style = _table_styles(prs)
    _background(canvas, slide)
    layout = slide.slide_layout
    # master elements first, then layout, then the slide itself
    _draw_shapes(canvas, layout.slide_master.shapes, skip_placeholders=True)
    _draw_shapes(canvas, layout.shapes, skip_placeholders=True)
    _draw_shapes(canvas, slide.shapes)
    return canvas.image


def render_slide(pptx, index=0, width=1920):
    """Render slide ``index`` (0-based) of a pptx path or Presentation as a PIL image."""
    prs = pptx if hasattr(pptx, "slides") else Presentation(pptx)
    return _render(prs, prs.slides[index], width)


def render_slides(prs, indexes=None, width=1920):
    """Render the slides at ``indexes`` (0-based; all slides if None) as a list of PIL images."""
    if not hasattr(prs, "slides"):
        prs = Presentation(prs)
    slides = list(prs.slides)
    if indexes is None:
        indexes = range(len(slides))
    return [_render(prs, slides[i], width) for i in indexes]


def render_presentation(pptx, output_dir, width=1920):
    """Save every slide as ``slide_<n>.png`` in ``output_dir``; returns the file paths."""
    prs = pptx if hasattr(pptx, "slides") else Presentation(pptx)
    os.makedirs(output_dir, exist_ok=True)
    paths = []
    for i, slide in enumerate(prs.slides, start=1):
        path = os.path.join(output_dir, f"slide_{i}.png")
        _render(prs, slide, width).save(path)
        paths.append(path)
    return paths
